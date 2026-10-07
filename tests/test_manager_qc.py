"""Offline QC: populations, histograms, fits, floods, reports, per-file streams (spec 003 T10, T33; spec 007 T11).

Moved from scripts/petsys_manager_qc_check.py. The reference outputs (scripts_cornell/cornell_system_validation.py)
are the golden files under tests/data/golden/qc/ (T8); no test loads the reference.
"""

import ast
from collections import defaultdict
import json
import os
from pathlib import Path
import tracemalloc
import unittest
from unittest.mock import patch

import numpy as np
import pytest
import yaml

from helpers import DATA
from manager_helpers import Geometry, PrivateOutput, QCFixtures
from src.cornell import qc, qc_report
from src.cornell.inputs import InputError, load_processing_config
from src.petsys_manager.contracts import InputDescriptor, SourceMode
from src.petsys_manager.settings import ProcessingLimits, ProfileError, RunOptions
from src.read_compact import read_binary_file

GOLDEN = DATA / "golden" / "qc"
LEGACY_HALVES = {2: [mm for mm in range(16) if mm not in qc.LEGACY_HALF_MINIMODULES]}


def pdf_lines(path):
    from pypdf import PdfReader
    return [line.strip() for page in PdfReader(str(path)).pages for line in page.extract_text().splitlines()
            if line.strip()]


def is_subsequence(needles, haystack):
    position = 0
    for needle in needles:
        while position < len(haystack) and haystack[position] != needle:
            position += 1
        if position == len(haystack):
            return False
        position += 1
    return True


def golden(name):
    """A reference output frozen at T8 (FR-3), parsed."""
    return json.loads((GOLDEN / name).read_bytes().decode("utf-8"))


def keyed(pairs):
    """Golden ``[key, value]`` pairs as a dict with tuple keys."""
    return {tuple(key) if isinstance(key, list) else key: value for key, value in pairs}


def energies(pairs):
    """Golden ``[key, {"dtype", "values"}]`` pairs as a dict of arrays."""
    return {key: np.array(value["values"], dtype=value["dtype"]) for key, value in keyed(pairs).items()}


def occupancy(reference):
    """The reference channel occupancy per SuperModule."""
    return {sm: dict(channels) for sm, channels in reference["occupancy"]}


@pytest.mark.fr("003-FR-2", "003-FR-9", "003-FR-14", "003-FR-15", "003-FR-16")  # spec 003 T10, T33
class QCChecks(QCFixtures, PrivateOutput, unittest.TestCase):
    fixture_prefix = "petsys-manager-qc-"

    def histogram(self, values):
        return np.histogram(np.array(values), bins=qc.PHOTOPEAK_BINS, range=qc.PHOTOPEAK_RANGE)[0]

    # Checks -----------------------------------------------------------------

    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_manual_pair_side_hit_counts_and_occupancy(self):
        descriptor = self.compact("manual.ldat", self.manual_records())
        result = self.qc_run([descriptor], plots=True, slabs=True)
        (sample,) = result.files
        self.assertEqual((sample.validated_records, sample.records_read, sample.pairs_processed), (5, 5, 5))
        self.assertEqual((sample.occupancy_pairs, sample.occupancy_hits), (3, 42))   # unresolved pair included
        self.assertEqual((sample.accepted_pairs, sample.accepted_sides, sample.stopped_at_limit), (2, 4, False))
        self.assertEqual(sample.rejected, {"min_channels": 1, "minimodule_channels": 1, "no_energy_minimodule": 0,
                                           "unresolved_slab": 1})
        self.assertEqual(sample.slab_flags, {"edge": 0, "single_time_random": 0, "non_adjacent": 1, "adjacent": 5})
        self.assertEqual(result.slab_counts, {(1, 0, 5): 2, (2, 0, 7): 2})
        fits = {e.key: e for e in result.minimodule_fits}
        self.assertEqual({k: e.samples for k, e in fits.items()}, {(1, 0): 2, (2, 0): 2})   # energy: resolved only
        self.assertAlmostEqual(fits[(1, 0)].sample_mean, 100.0, places=9)                  # 0.1 a.u. hit cut
        occupancy_hits = sum(sum(channels.values()) for channels in result.occupancy.values())
        self.assertEqual(occupancy_hits, 42)
        self.assertIn(self.geometry.time[(1, 1)][3][0], result.occupancy[1])               # unresolved side hits
        reference = golden("manual.json")
        self.assertEqual(reference["accepted_pairs"], 2)
        self.assertEqual(occupancy(reference), result.occupancy)
        self.assertEqual(keyed(reference["slab_counts"]), result.slab_counts)
        self.assertEqual({k: len(v) for k, v in energies(reference["minimodule_energy"]).items()}, {(1, 0): 2, (2, 0): 2})
        no_cut = self.qc_run([descriptor], self.processing({"en_min_ch": 0.0}, "nocut.yaml"), plots=True)
        self.assertEqual(no_cut.files[0].occupancy_hits, 43)
        self.assertAlmostEqual({e.key: e for e in no_cut.minimodule_fits}[(1, 0)].sample_mean, 100.05, places=6)

    def test_qc_reader_equals_reference_reader(self):
        descriptor = self.inputs(count=300, names=("reader.ldat",))[0]
        with open(os.devnull, "w") as quiet, patch("sys.stderr", quiet):
            reference = list(read_binary_file(str(descriptor.path), 0.2, group_events=False))
        ours = list(qc.read_pairs(descriptor.path, 0.2))
        self.assertEqual([tuple(map(list, pair)) for pair in reference], [tuple(map(list, pair)) for pair in ours])
        self.assertTrue(any(hit[1] < 0.2 for pair in qc.read_pairs(descriptor.path, 0.0) for side in pair
                            for hit in side))

    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_histograms_fits_occupancy_floods_match_reference(self):
        descriptors = self.inputs()
        result = self.qc_run(descriptors, plots=True, slabs=True, flush_sides=97)
        reference = golden("inputs_1500.json")
        totals = result.totals()
        self.assertEqual(totals["accepted_pairs"], reference["accepted_pairs"])
        self.assertEqual(occupancy(reference), result.occupancy)
        self.assertEqual(keyed(reference["slab_counts"]), result.slab_counts)
        self.assertGreater(totals["slab_flags"]["single_time_random"], 0)
        for reason in ("min_channels", "minimodule_channels", "unresolved_slab"):
            self.assertGreater(totals["rejected"][reason], 0, reason)
        statuses = defaultdict(int)
        for entries, label in ((result.minimodule_fits, "minimodule"), (result.slab_fits, "slab")):
            energy, fits = energies(reference[f"{label}_energy"]), keyed(reference["fits"][label])
            self.assertEqual(sorted(e.key for e in entries), sorted(energy))
            for entry in entries:
                values = np.array(energy[entry.key])
                np.testing.assert_array_equal(entry.histogram, self.histogram(values))
                fit = fits[entry.key]          # the reference fit_photopeak_worker result
                mu, sigma, res, error = fit["mu"], fit["sigma"], fit["resolution_percent"], fit["fallback"]
                statuses[entry.status] += 1
                self.assertEqual(entry.samples, len(values))
                if error:
                    self.assertFalse(entry.available)
                    self.assertIsNone(entry.mu)
                    self.assertAlmostEqual(entry.sample_mean, mu, delta=1e-9 * abs(mu))
                    self.assertAlmostEqual(entry.sample_std, sigma, delta=1e-9 * abs(sigma) + 1e-12)
                else:
                    self.assertEqual(entry.status, "fitted", entry)
                    self.assertEqual((entry.mu, abs(entry.sigma)), (mu, abs(sigma)))
                    self.assertAlmostEqual(entry.resolution_percent, abs(res), delta=1e-12)
        self.assertGreater(statuses["fitted"], 0)
        self.assertGreater(statuses["sparse"], 0)
        floods = energies(reference["flood_points"])
        for sm, counts in result.floods.items():
            points = [p for key, value in floods.items() if key[0] == sm for p in value
                      if qc.FLOOD_ENERGY[0] <= p[2] <= qc.FLOOD_ENERGY[1]]
            expected, _, _ = np.histogram2d([p[0] for p in points], [p[1] for p in points], bins=(500, 500),
                                            range=[[0, 105], [0, 105]])
            np.testing.assert_array_equal(counts, expected)
        self.assertEqual(sorted(result.floods), sorted({key[0] for key in floods}))
        excluded = sum(1 for value in floods.values() for p in value
                       if not qc.FLOOD_ENERGY[0] <= p[2] <= qc.FLOOD_ENERGY[1])
        self.assertGreater(excluded, 0)

    @staticmethod
    def fingerprint(result):
        """Everything a QC report is built from (T33 equality across worker counts)."""
        fits = [(e.key, e.samples, e.in_histogram, e.status, e.mu, e.sigma, e.resolution_percent, e.sample_mean,
                 e.sample_std, e.histogram.tobytes()) for e in result.minimodule_fits + result.slab_fits]
        return (result.files, result.occupancy, result.slab_counts, fits,
                {k: v.tobytes() for k, v in result.floods.items()}, result.findings, result.storage_bytes)

    @pytest.mark.slow  # ~7 s
    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_parallel_files_with_per_file_seeds(self):
        """FR-15 (T33): with a QC seed each file draws from its own Python random stream into its own
        accumulators, merged in input order: 1, 2 and 3 workers and repeated runs give the same result; the
        reference seeded the same way per file gives the same populations and histograms; another seed changes
        only slab-draw dependent values; seeds and workers are recorded; the unseeded path is unchanged."""
        descriptors = self.inputs(count=900, names=("par_coinc_1.ldat", "par_coinc_2.ldat", "par_coinc_3.ldat"))
        options = dict(plots=True, slabs=True, flush_sides=97)
        base = self.qc_run(descriptors, qc_seed=5, **options)
        self.assertGreater(base.totals()["slab_flags"]["single_time_random"], 0)
        for workers in (2, 3, 8):
            other = self.qc_run(descriptors, qc_seed=5, workers=workers, **options)
            self.assertEqual(self.fingerprint(other), self.fingerprint(base), workers)
            self.assertEqual(other.workers, min(workers, 3))
        again = self.qc_run(descriptors, seed=99, qc_seed=5, **options)        # the global seed does not matter
        self.assertEqual(self.fingerprint(again), self.fingerprint(base))
        seeds = [qc.file_seed(5, index) for index in range(3)]
        self.assertEqual(base.random_streams["file_seeds"], seeds)
        self.assertEqual((base.random_streams["qc_seed"], base.workers), (5, 1))
        reference = golden("per_file_seeds.json")
        self.assertEqual(base.totals()["accepted_pairs"], reference["accepted_pairs"])
        self.assertEqual(occupancy(reference), base.occupancy)
        self.assertEqual(keyed(reference["slab_counts"]), base.slab_counts)
        for entries, energy in ((base.minimodule_fits, energies(reference["minimodule_energy"])),
                                (base.slab_fits, energies(reference["slab_energy"]))):
            self.assertEqual(sorted(e.key for e in entries), sorted(energy))
            for entry in entries:
                np.testing.assert_array_equal(entry.histogram, self.histogram(energy[entry.key]))
                self.assertAlmostEqual(entry.sample_mean, float(np.mean(energy[entry.key])),
                                       delta=1e-9 * abs(entry.sample_mean))
        other = self.qc_run(descriptors, qc_seed=6, **options)
        self.assertNotEqual(other.slab_counts, base.slab_counts)
        self.assertEqual(other.occupancy, base.occupancy)                         # before the slab draw
        self.assertEqual([f.records_read for f in other.files], [f.records_read for f in base.files])
        summary = qc_report.summary(base, [])
        self.assertEqual(summary["sampling"]["random_streams"]["file_seeds"], seeds)
        self.assertEqual(summary["sampling"]["workers"], 1)
        unseeded = self.qc_run(descriptors, **options)
        self.assertIsNone(unseeded.random_streams["qc_seed"])
        self.assertEqual(qc_report.summary(unseeded, [])["sampling"]["workers"], 1)
        for bad in ({"workers": 2}, {"qc_seed": -1}, {"qc_seed": 1.5}, {"qc_seed": 5, "workers": 0}):
            with self.subTest(bad=bad), self.assertRaises(InputError):
                self.qc_run(descriptors, **bad)
        calls = []
        with self.assertRaises(qc.QCCancelled):
            self.qc_run(descriptors, qc_seed=5, workers=2, cancelled=lambda: calls.append(1) or len(calls) > 2)
        progress = []
        self.qc_run(descriptors, qc_seed=5, workers=2, progress=lambda path, records: progress.append(path))
        self.assertEqual(progress, [d.path for d in descriptors])

    @pytest.mark.fr("003-FR-24", "007-FR-3")  # FR-24; golden files
    def test_qc_per_file_stopping_before_at_after_limit(self):
        records = self.manual_records()
        valid = [records[0], records[1]] * 3                       # 6 accepted pairs
        descriptor = self.compact("limit.ldat", valid + [records[3], records[2]])
        cases = {5: (6, 5, True), 6: (7, 6, True), 7: (8, 6, False), 100: (8, 6, False)}
        for limit, (read, accepted, stopped) in cases.items():
            with self.subTest(limit=limit):
                sample = self.qc_run([descriptor], pair_limit=limit).files[0]
                self.assertEqual((sample.records_read, sample.accepted_pairs, sample.stopped_at_limit),
                                 (read, accepted, stopped))
                self.assertEqual(sample.validated_records, read)            # FR-24: the records read
                self.assertEqual(sample.records_in_file, None if stopped else 8)
                self.assertEqual(sample.pairs_processed, read - stopped)
        self.assertEqual(qc.PAIR_LIMIT, ProcessingLimits().qc_pair_limit)
        self.assertEqual(qc.PAIR_LIMIT, 1_000_001)
        # Reference counter seeded at 999,998 / 999,999 / 1,000,000 equals our limit 3 / 2 / 1 per file.
        cases = golden("pair_limit.json")["cases"]
        for case, (start, limit) in zip(cases, ((999_998, 3), (999_999, 2), (1_000_000, 1)), strict=True):
            with self.subTest(start=start):
                self.assertEqual(case["counter_start"], start)
                sample = self.qc_run([descriptor], pair_limit=limit).files[0]
                self.assertEqual((case["accepted_pairs"], case["records_read"]), (sample.accepted_pairs, sample.records_read))
                self.assertTrue(sample.stopped_at_limit)

    @pytest.mark.slow  # ~7 s
    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_option_combinations_and_reference_output_types(self):
        descriptors = self.inputs(count=800)
        reference = golden("options.json")
        legacy_files = {(o["plots"], o["slabs"]): set(o["files"]) for o in reference["outputs"]}
        with self.assertRaises(InputError):
            self.qc_run(descriptors, plots=False, slabs=True)
        with self.assertRaises(ProfileError):
            RunOptions(plots=False, slabs=True)
        for plots, slabs in ((False, False), (True, False), (True, True)):
            with self.subTest(plots=plots, slabs=slabs):
                result = self.qc_run(descriptors, plots=plots, slabs=slabs)
                self.assertEqual(result.totals()["accepted_pairs"], reference["accepted_pairs"])
                ours = self.results / f"ours_{plots}_{slabs}"
                written = qc_report.write_report(result, ours)
                self.assertEqual(written[-1].name, qc_report.SUMMARY)
                names = {p.name for p in ours.iterdir()} - {qc_report.SUMMARY}
                legacy_names = legacy_files[(plots, slabs)]
                # Documented difference: slab distributions only for the selected map's SuperModules (not SM 0-29).
                selected = {f"slab_distribution_SM{sm}{suffix}.png" for sm in (1, 2) for suffix in ("", "_no_data")}
                legacy_names = {n for n in legacy_names if not n.startswith("slab_distribution_SM") or n in selected}
                self.assertEqual(names, legacy_names)
                self.assertEqual(bool(result.minimodule_fits), plots)
                self.assertEqual(bool(result.slab_fits), slabs)
                content = json.loads((ours / qc_report.SUMMARY).read_text(encoding="utf-8"))
                self.assertEqual(content["options"], {"plots": plots, "slabs": slabs})
                self.assertEqual({o["path"] for o in content["outputs"]}, names)
        with self.assertRaises(FileExistsError):
            qc_report.write_report(result, ours)

    def test_qc_source_metadata_distinct_and_numbers_identical(self):
        descriptors = self.inputs(count=400)
        seen = {}
        for mode, duration, text in ((SourceMode.WITH, 60, "with source, 60 s acquisition"),
                                     (SourceMode.WITHOUT, 180, "without source, 180 s acquisition"),
                                     (None, None, "not recorded (offline analysis)")):
            result = self.qc_run(descriptors, source_mode=mode, acquisition_time_s=duration)
            directory = self.results / f"source_{mode}"
            qc_report.write_report(result, directory)
            content = json.loads((directory / qc_report.SUMMARY).read_text(encoding="utf-8"))
            self.assertEqual((content["source_mode"], content["acquisition_time_s"]),
                             (None if mode is None else mode.value, None if duration is None else float(duration)))
            self.assertIn(f"Source mode: {text}", " ".join(pdf_lines(directory / qc_report.PDF)))
            seen[mode] = (content["totals"], content["findings"])
        self.assertEqual(len({json.dumps(v, sort_keys=True) for v in seen.values()}), 1)
        with self.assertRaises(InputError):
            self.qc_run(descriptors, source_mode="with")

    @pytest.mark.slow  # ~5 s
    def test_qc_unavailable_fits_and_raw_units(self):
        descriptors = self.inputs(count=1500)
        result = self.qc_run(descriptors, plots=True, slabs=True)
        directory = self.results / "fits"
        qc_report.write_report(result, directory)
        import openpyxl
        sheet = openpyxl.load_workbook(directory / qc_report.EXCEL)["Photopeak Values"]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], qc_report.EXCEL_COLUMNS)
        statuses = {(r[0], r[1]): r for r in rows[1:]}
        unavailable = [e for e in result.minimodule_fits if not e.available]
        self.assertTrue(unavailable)
        for entry in unavailable:
            row = statuses[entry.key]
            self.assertEqual(row[2:5], (None, None, None))
            self.assertEqual((row[5], row[6]), (entry.status, entry.samples))
        provenance = dict((r[0], r[1]) for r in openpyxl.load_workbook(directory / qc_report.EXCEL)["Provenance"]
                          .iter_rows(values_only=True))
        self.assertTrue(provenance["Units"].startswith("a.u."))
        content = json.loads((directory / qc_report.SUMMARY).read_text(encoding="utf-8"))
        self.assertTrue(content["energy_units"].startswith("a.u."))
        self.assertIsNone(content["calibration"])
        listed = {tuple(e["key"]): e for e in content["minimodule_fits"]["entries"]}
        for entry in unavailable:
            self.assertIsNone(listed[entry.key]["mu_au"])
        self.assertTrue(content["slab_fits"]["non_fitted"])
        self.assertTrue(all(e["status"] != "fitted" or e["message"] for e in content["slab_fits"]["non_fitted"]))
        # Forced fit failures, errors and non-physical results are unavailable too.
        accumulator = qc.EnergyHistograms()
        accumulator.add([(9, 9)] * 1000, list(np.random.default_rng(1).normal(100, 10, 1000)))
        cases = [(RuntimeError("Fitting failed: forced"), "fit_failed"), (ValueError("forced"), "fit_error"),
                 ((None, None, [1.0, -5.0, 3.0], None, None), "invalid_result"),
                 ((None, None, [1.0, 100.0, -9.0], None, None), "fitted")]
        for outcome, status in cases:
            effect = dict(side_effect=outcome) if isinstance(outcome, Exception) else dict(return_value=outcome)
            with self.subTest(status=status), patch.object(qc, "fit_gaussian", **effect):
                (entry,) = qc.fit_photopeaks(accumulator, qc.MINIMODULE_CB)
                self.assertEqual(entry.status, status)
                self.assertEqual(entry.available, status == "fitted")
                if status == "fitted":
                    self.assertEqual(entry.sigma, 9.0)
                    self.assertIn("absolute value", entry.message)

    @pytest.mark.slow  # ~5 s
    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_pdf_excel_and_plot_content_match_reference(self):
        descriptors = self.inputs(count=1500)
        config = self.processing({"en_min_ch": 0.2, "unpopulated_minimodules": LEGACY_HALVES}, "legacy.yaml")
        result = self.qc_run(descriptors, config, plots=True, slabs=True)
        ours = self.results / "content"
        qc_report.write_report(result, ours)
        reference = golden("report_content.json")
        self.assertEqual(result.expected, dict(reference["expected_channels"]))
        self.assertEqual(result.findings.missing, dict(reference["missing_channels"]))
        legacy_lines = reference["pdf_lines"]                # without the "Report generated" line
        self.assertTrue(is_subsequence(legacy_lines, pdf_lines(ours / qc_report.PDF)))
        self.assertIn("NOT OK", " ".join(legacy_lines))
        import openpyxl
        legacy_rows = [tuple(row) for row in reference["excel_rows"]]
        our_rows = list(openpyxl.load_workbook(ours / qc_report.EXCEL).active.iter_rows(values_only=True))
        self.assertEqual(our_rows[0][:5], legacy_rows[0])
        fitted = {e.key for e in result.minimodule_fits if e.available}
        self.assertEqual([r[:5] for r in our_rows[1:] if (r[0], r[1]) in fitted],
                         [r for r in legacy_rows[1:] if (r[0], r[1]) in fitted])
        for row in legacy_rows[1:]:
            if (row[0], row[1]) not in fitted:       # reference fallback mean/std written as Mu/Sigma
                ours_row = next(r for r in our_rows if (r[0], r[1]) == (row[0], row[1]))
                self.assertAlmostEqual(ours_row[7], row[2], delta=1e-9 * abs(row[2]))
        self.assertEqual(sorted(map(tuple, reference["photopeak_keys"])), [e.key for e in result.minimodule_fits])
        content = json.loads((ours / qc_report.SUMMARY).read_text(encoding="utf-8"))
        self.assertEqual(content["expectation"]["expected_now_skipped_by_legacy"], [])
        self.assertEqual(content["expectation"]["skipped_now_expected_by_legacy"], [])
        for output in content["outputs"]:
            self.assertEqual(len(output["sha256"]), 64)

    def test_qc_changed_inactive_module_expectation_is_explained(self):
        descriptors = self.inputs(count=600)
        result = self.qc_run(descriptors)
        legacy = {(2, mm) for mm in range(16) if mm not in qc.LEGACY_HALF_MINIMODULES}
        self.assertEqual(result.legacy_absent, legacy)
        self.assertEqual(result.absent, frozenset())
        missing = set(result.findings.missing_minimodules)
        self.assertTrue(legacy <= missing)               # no sensors there: now reported as not observed
        directory = self.results / "expectation"
        qc_report.write_report(result, directory)
        content = json.loads((directory / qc_report.SUMMARY).read_text(encoding="utf-8"))
        self.assertEqual({tuple(k) for k in content["expectation"]["expected_now_skipped_by_legacy"]}, legacy)
        text = " ".join(pdf_lines(directory / qc_report.PDF))
        self.assertIn("legacy hardcoded half-SuperModule rule is not applied; 8 minimodules", text)
        declared = self.processing({"en_min_ch": 0.2, "unpopulated_minimodules": {1: [0], **LEGACY_HALVES}}, "d.yaml")
        result = self.qc_run(descriptors, declared)
        self.assertEqual(set(result.findings.missing_minimodules) & legacy, set())
        self.assertIn((1, 0), result.findings.unexpected_minimodules)   # declared absent, hits observed
        self.assertNotIn((1, 0), result.findings.expected_minimodules)

    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_bounded_storage_and_incremental_merge(self):
        one = self.inputs(count=600, names=("one.ldat",))
        many = [self.compact(f"copy_{i}.ldat", self.records(600, seed=7)) for i in range(4)]
        self.qc_run(one, plots=True, slabs=True, flush_sides=64)        # warm imports/caches
        peaks, results = [], []
        for descriptors in (one, many):
            tracemalloc.start()
            results.append(self.qc_run(descriptors, plots=True, slabs=True, flush_sides=64))
            peaks.append(tracemalloc.get_traced_memory()[1])
            tracemalloc.stop()
        single, quadruple = results
        self.assertEqual(single.storage_bytes, quadruple.storage_bytes)
        self.assertLess(peaks[1] - peaks[0], 256 * 1024, peaks)           # 3 more files, no list growth
        for a, b in zip(single.minimodule_fits, quadruple.minimodule_fits):
            self.assertEqual((a.key, 4 * a.samples), (b.key, b.samples))
            np.testing.assert_array_equal(4 * a.histogram, b.histogram)
        self.assertEqual(4 * sum(single.slab_counts.values()), sum(quadruple.slab_counts.values()))
        self.assertEqual(single.floods.keys(), quadruple.floods.keys())
        for sm in single.floods:
            self.assertEqual(4 * int(single.floods[sm].sum()), int(quadruple.floods[sm].sum()))
        self.assertGreater(golden("bounded.json")["flood_sides"], 4 * 600)   # reference keeps every side

    @pytest.mark.fr("007-FR-3")  # golden files
    def test_qc_reference_failure_case_becomes_counted_rejection(self):
        zero = self.processing({"en_min_ch": 0.0}, "zero.yaml")
        times = self.geometry.time[(1, 0)]
        side = [(1, 12.0, times[2][0]), (2, 5.0, times[3][0])]
        side += [(3, 0.0, self.geometry.energy[(1, 0)][i][0]) for i in range(3)]
        side += [(4, 0.0, self.geometry.energy[(1, 1)][i][0]) for i in range(2)]   # two minimodules, no energy
        records = self.manual_records()
        descriptor = self.compact("zero_energy.ldat", [(side, records[0][1]), records[0]])
        result = self.qc_run([descriptor], zero)
        self.assertEqual(result.files[0].rejected["no_energy_minimodule"], 1)
        self.assertEqual(result.files[0].accepted_pairs, 1)
        self.assertEqual(golden("zero_energy.json")["reference_raises"]["exception"], "TypeError")

    def test_qc_generalized_layout_for_selected_map(self):
        self.write_map({7: [0, 0, 0], 41: [2, 0, 0]})
        config = self.processing({"en_min_ch": 0.2}, "wide.yaml")
        self.geometry = Geometry(config.mapping)
        with patch("manager_helpers.QC_SPECS", [((7, 0), 2, +1, 100.0, 9.0, 50), ((41, 3), 3, +1, 80.0, 7.0, 50)]):
            descriptors = self.inputs(count=400)
        result = self.qc_run(descriptors, config, plots=True)
        directory = self.results / "wide"
        qc_report.write_report(result, directory)
        names = {p.name for p in directory.iterdir()}
        self.assertTrue({"slab_distribution_SM7.png", "slab_distribution_SM41.png", "floodmap_SM_7.png",
                         "floodmap_SM_41.png", "floodmap_all_SM.png", "photopeak_SM_7.png"} <= names)
        self.assertFalse(any(n.startswith("slab_distribution_SM0") for n in names))
        self.assertEqual(qc_report.cassettes(result.expected), {2: [7], 13: [41]})
        self.assertEqual(result.legacy_absent, {(41, mm) for mm in range(16) if mm not in qc.LEGACY_HALF_MINIMODULES})

    @pytest.mark.fr("003-FR-24")  # FR-24
    def test_qc_input_contract_rejections_and_cancellation(self):
        descriptors = self.inputs(count=200)
        fixed = InputDescriptor(descriptors[0].path, "fixed", "coincidence")
        cases = {
            "fixed": dict(descriptors=[fixed]), "empty": dict(descriptors=[]),
            "duplicate": dict(descriptors=[descriptors[0], descriptors[0]]),
            "untyped": dict(descriptors=[str(descriptors[0].path)]),
            "slabs_only": dict(slabs=True), "limit": dict(pair_limit=0), "bool_limit": dict(pair_limit=True),
            "duration": dict(acquisition_time_s=-1), "plots_type": dict(plots=1),
        }
        for name, options in cases.items():
            with self.subTest(name=name), self.assertRaises(InputError):
                self.qc_run(options.pop("descriptors", descriptors), **options)
        bad = self.compact("unmapped.ldat", [([(1, 50.0, 999_999)] + self.manual_records()[0][0],
                                              self.manual_records()[0][1])])
        with self.assertRaises(InputError):
            self.qc_run([bad])
        truncated = self.root / "truncated.ldat"
        truncated.write_bytes(descriptors[0].path.read_bytes()[:-3])
        with self.assertRaises(InputError):
            self.qc_run([InputDescriptor(truncated, "compact", "coincidence")])
        missing_cut = self.root / "configs/nocut.yaml"
        missing_cut.write_text(yaml.safe_dump({"map_file": "maps/selected.yaml", "min_ch": 4}), encoding="utf-8")
        with self.assertRaises(InputError):
            load_processing_config(missing_cut, processing_root=self.root, action="qc_analyze")
        calls = []
        with self.assertRaises(qc.QCCancelled):
            # FR-24: no validation pass polls; polls are per file and every 1024 records (2 for these files).
            qc.run_qc(descriptors, self.config, cancelled=lambda: calls.append(1) or len(calls) > 1)
        with self.assertRaises(InputError):
            qc_report.write_report(self.qc_run(descriptors), self.root / "absent_parent/results")
        self.assertFalse((self.root / "absent_parent").exists())
        stamp = qc_report.default_directory(self.results, __import__("datetime").datetime(2026, 10, 1, 9, 8, 7))
        self.assertEqual(stamp, self.results / "20261001-090807")

    def test_qc_tracked_modules_never_import_local_scripts(self):
        forbidden = {"scripts", "scripts_cornell", "scripts_gui", "docopt", "multiprocessing", "tkinter",
                     "customtkinter", "pandas", "natsort", "colorama", "tqdm", "gui"}
        for module in (qc, qc_report):
            tree = ast.parse(Path(module.__file__).read_text(encoding="utf-8"))
            names = set()
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    names |= {alias.name.split(".")[0] for alias in node.names}
                elif isinstance(node, ast.ImportFrom) and node.module and node.level == 0:
                    names.add(node.module.split(".")[0])
            self.assertFalse(names & forbidden, (module.__name__, names & forbidden))
        for module in (qc, qc_report):
            top = set()
            for node in ast.parse(Path(module.__file__).read_text(encoding="utf-8")).body:
                if isinstance(node, ast.Import):
                    top |= {alias.name.split(".")[0] for alias in node.names}
                elif isinstance(node, ast.ImportFrom) and node.module:
                    top.add(node.module.split(".")[0])
            self.assertFalse(top & {"reportlab", "openpyxl", "matplotlib"}, module.__name__)   # lazy, headless
